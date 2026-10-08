def clear_terminal():
    print("\033[2J\033[H", end='')


def print_color(text: str = "", color: str = "reset"):
    colors = {
        "red": "\033[31m",
        "green": "\033[32m",
        "yellow": "\033[33m",
        "blue": "\033[34m",
        "reset": "\033[0m"
    }
    print(f"{colors.get(color, colors['reset'])}{text}{colors['reset']}")


def print_dict(d):
    import json
    print(json.dumps(d, indent=4, sort_keys=True)) 


def _adjust_excel_column_widths(file_path: str):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    wb = load_workbook(file_path)
    for ws in wb.worksheets:
        for col in ws.columns:
            max_length = 0
            col_letter = get_column_letter(col[0].column)

            for cell in col:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass

            ws.column_dimensions[col_letter].width = max_length + 2
    wb.save(file_path)


def to_excel(df, file_path: str, adjust: bool = True):
    df.to_excel(file_path, index=False)
    if adjust:
        _adjust_excel_column_widths(file_path)


def to_csv(df, file_path: str):
    df.to_csv(file_path, index=False)


def save_dataframes_to_excel(
    dataframes: list,
    sheet_names: list[str],
    file_path: str,
    adjust: bool = True
):
    import pandas as pd

    if len(dataframes) != len(sheet_names):
        raise ValueError("dataframes and sheet_names must have the same length")

    with pd.ExcelWriter(file_path, engine='xlsxwriter') as writer:
        for df, name in zip(dataframes, sheet_names):
            df.to_excel(writer, sheet_name=name, index=False)
    if adjust:
        _adjust_excel_column_widths(file_path)


def is_valid_language(lang: str) -> bool:
    import pycountry
    try:
        pycountry.languages.lookup(lang)
        return True
    except LookupError:
        return False